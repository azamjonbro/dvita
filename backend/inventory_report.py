"""Ombordagi tovarlar qiymati hisoboti (tannarx, retail, potensial yalpi foyda).

Bazaga tegmaydigan sof hisob-kitob va Excel eksporti — unit-test qilish oson.

Formulalar:
    jami tannarx        = Σ qoldiq × tannarx
    jami retail         = Σ qoldiq × amaldagi sotuv narxi (chegirmasiz)
    potensial yalpi foyda = retail − tannarx
    ustama %            = foyda ÷ tannarx × 100
    marja %             = foyda ÷ retail × 100

Bu SOF foyda emas: arenda, oylik, reklama, soliq va boshqa xarajatlar ayirilmagan.
"""
from __future__ import annotations

import io
from typing import Dict, Iterable, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NO_BRANCH_KEY = "none"
NO_BRANCH_NAME = "Filialsiz / umumiy"


def _pct(part: float, whole: float) -> Optional[float]:
    return round(part / whole * 100, 2) if whole else None


def _num(value) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def item_row(product: dict, branch_name: str) -> Optional[dict]:
    """Bitta mahsulot qatori; qoldig'i yo'q mahsulot hisobotga kirmaydi."""
    qty = max(int(_num(product.get("stock"))), 0)
    if qty <= 0:
        return None
    unit_cost = _num(product.get("cost_price"))
    unit_price = _num(product.get("price"))
    cost_total = round(qty * unit_cost, 2)
    retail_total = round(qty * unit_price, 2)
    profit = round(retail_total - cost_total, 2)
    return {
        "product_id": product.get("product_id") or product.get("id"),
        "name": product.get("name", ""),
        "sku": product.get("sku") or "",
        "barcode": product.get("barcode") or "",
        "branch_id": product.get("branch_id") or NO_BRANCH_KEY,
        "branch_name": branch_name,
        "qty": qty,
        "unit_cost": unit_cost,
        "unit_price": unit_price,
        "cost_total": cost_total,
        "retail_total": retail_total,
        "profit": profit,
        "margin_pct": _pct(profit, retail_total),
    }


def summarize(rows: Iterable[dict]) -> dict:
    rows = list(rows)
    cost = round(sum(r["cost_total"] for r in rows), 2)
    retail = round(sum(r["retail_total"] for r in rows), 2)
    profit = round(retail - cost, 2)
    return {
        "sku_count": len(rows),
        "total_qty": sum(r["qty"] for r in rows),
        "cost_total": cost,
        "retail_total": retail,
        "profit": profit,
        "markup_pct": _pct(profit, cost),
        "margin_pct": _pct(profit, retail),
    }


def build_report(products_by_branch: Dict[str, List[dict]], branch_names: Dict[str, str]) -> dict:
    """products_by_branch: {branch_key: [mahsulot, ...]} -> filiallar, umumiy va mahsulotlar."""
    branches, items = [], []
    for key, products in products_by_branch.items():
        name = branch_names.get(key, NO_BRANCH_NAME)
        rows = [r for r in (item_row(p, name) for p in products) if r]
        rows.sort(key=lambda r: r["name"].lower())
        items.extend(rows)
        branches.append({"branch_id": key, "branch_name": name, **summarize(rows)})
    branches.sort(key=lambda b: (b["branch_id"] == NO_BRANCH_KEY, b["branch_name"].lower()))
    return {"branches": branches, "total": summarize(items), "items": items}


# ---------- Excel ----------
_HEAD_FILL = PatternFill("solid", fgColor="1F3A33")
_HEAD_FONT = Font(bold=True, color="FFFFFF")
_TOTAL_FONT = Font(bold=True)
_MONEY = '#,##0'
_PCT = '0.0"%"'


def _write_table(ws, headers: List[str], rows: List[list], money_cols=(), pct_cols=(), total_row=None):
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.fill, cell.font = _HEAD_FILL, _HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    if total_row:
        ws.append(total_row)
        for cell in ws[ws.max_row]:
            cell.font = _TOTAL_FONT
    for col in range(1, len(headers) + 1):
        letter = get_column_letter(col)
        for cell in ws[letter][1:]:
            if col in money_cols:
                cell.number_format = _MONEY
            elif col in pct_cols:
                cell.number_format = _PCT
        width = max(len(str(c.value or "")) for c in ws[letter])
        ws.column_dimensions[letter].width = min(max(12, width + 2), 60)
    ws.freeze_panes = "A2"


def build_xlsx(report: dict, date_label: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Filiallar"
    ws.append([f"Ombordagi tovarlar qiymati — {date_label}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Potensial yalpi foyda — sof foyda emas: arenda, oylik, reklama, soliq va boshqa xarajatlar ayirilmagan."])
    ws.append([])
    t = report["total"]
    branch_rows = [[b["branch_name"], b["sku_count"], b["total_qty"], b["cost_total"], b["retail_total"],
                    b["profit"], b["markup_pct"], b["margin_pct"]] for b in report["branches"]]
    header_ws_offset = ws.max_row
    _write_table(
        ws,
        ["Filial", "Mahsulot turi", "Qoldiq (dona)", "Jami tannarx", "Jami retail", "Potensial yalpi foyda",
         "Ustama %", "Marja %"],
        branch_rows,
        money_cols=(4, 5, 6), pct_cols=(7, 8),
        total_row=["Umumiy", t["sku_count"], t["total_qty"], t["cost_total"], t["retail_total"], t["profit"],
                   t["markup_pct"], t["margin_pct"]],
    )
    ws.freeze_panes = f"A{header_ws_offset + 2}"

    ws2 = wb.create_sheet("Mahsulotlar")
    _write_table(
        ws2,
        ["Mahsulot", "SKU / artikul", "Shtrix-kod", "Filial", "Qoldiq", "Dona tannarxi", "Dona retail narxi",
         "Jami tannarx", "Jami retail", "Potensial foyda", "Marja %"],
        [[r["name"], r["sku"], r["barcode"], r["branch_name"], r["qty"], r["unit_cost"], r["unit_price"],
          r["cost_total"], r["retail_total"], r["profit"], r["margin_pct"]] for r in report["items"]],
        money_cols=(6, 7, 8, 9, 10), pct_cols=(11,),
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
