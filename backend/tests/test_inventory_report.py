"""Ombordagi tovarlar qiymati hisoboti testlari."""
import io
import sys
import uuid
from pathlib import Path

import openpyxl
import pytest
import requests

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from inventory_report import build_report, build_xlsx  # noqa: E402

from .conftest import BASE_URL, DIRECTOR_EMAIL, DIRECTOR_PASSWORD  # noqa: E402


def test_formula_matches_spec_example():
    # Tannarxda 100 mln, retail 180 mln -> foyda 80 mln, ustama 80%, marja 44.4%
    report = build_report(
        {"b1": [{"id": "p1", "name": "A", "stock": 1000, "cost_price": 100_000, "price": 180_000}]},
        {"b1": "Namangan filiali"},
    )
    t = report["total"]
    assert (t["cost_total"], t["retail_total"], t["profit"]) == (100_000_000, 180_000_000, 80_000_000)
    assert t["markup_pct"] == 80.0 and t["margin_pct"] == 44.44


def test_multi_branch_rows_total_and_zero_stock_excluded():
    report = build_report(
        {
            "b1": [{"id": "1", "name": "X", "stock": 10, "cost_price": 10, "price": 18},
                   {"id": "2", "name": "Y", "stock": 0, "cost_price": 5, "price": 9}],      # qoldiqsiz — kirmaydi
            "b2": [{"id": "3", "name": "Z", "stock": 5, "cost_price": 50, "price": 86}],
        },
        {"b1": "Namangan", "b2": "Toshkent"},
    )
    rows = {b["branch_name"]: b for b in report["branches"]}
    assert (rows["Namangan"]["cost_total"], rows["Namangan"]["retail_total"], rows["Namangan"]["profit"]) == (100, 180, 80)
    assert (rows["Toshkent"]["cost_total"], rows["Toshkent"]["retail_total"], rows["Toshkent"]["profit"]) == (250, 430, 180)
    assert (report["total"]["cost_total"], report["total"]["retail_total"], report["total"]["profit"]) == (350, 610, 260)
    assert report["total"]["total_qty"] == 15 and report["total"]["sku_count"] == 2
    assert {i["name"] for i in report["items"]} == {"X", "Z"}
    item = next(i for i in report["items"] if i["name"] == "Z")
    assert (item["unit_cost"], item["unit_price"], item["profit"], item["margin_pct"]) == (50, 86, 180, 41.86)
    # Excel: filiallar varag'ida umumiy qator, mahsulotlar varag'ida detalizatsiya
    wb = openpyxl.load_workbook(io.BytesIO(build_xlsx(report, "2026-10-06")))
    summary = [r for r in wb["Filiallar"].iter_rows(values_only=True) if r and r[0] == "Umumiy"]
    assert summary and summary[0][3:6] == (350, 610, 260)
    assert wb["Mahsulotlar"].max_row == 3


def test_zero_cost_does_not_divide_by_zero():
    report = build_report({"b": [{"id": "1", "name": "Gift", "stock": 2, "cost_price": 0, "price": 0}]}, {"b": "B"})
    assert report["total"]["markup_pct"] is None and report["total"]["margin_pct"] is None


@pytest.fixture()
def director():
    s = requests.Session()
    r = s.post(f"{BASE_URL}/api/auth/login", json={"email": DIRECTOR_EMAIL, "password": DIRECTOR_PASSWORD}, timeout=20)
    assert r.status_code == 200, r.text
    s.headers.update({"Authorization": f"Bearer {r.json()['token']}"})
    return s


def test_inventory_value_api_and_excel(director):
    tag = uuid.uuid4().hex[:4]
    b1 = director.post(f"{BASE_URL}/api/branches", json={"name": f"Inv A {tag}", "code": f"IA-{tag}"}, timeout=20).json()["id"]
    b2 = director.post(f"{BASE_URL}/api/branches", json={"name": f"Inv B {tag}", "code": f"IB-{tag}"}, timeout=20).json()["id"]
    pids = []
    for branch, cost, price, stock in ((b1, 100, 180, 10), (b2, 250, 430, 2), (b2, 1, 2, 0)):
        r = director.post(f"{BASE_URL}/api/products", json={
            "name": f"Inv {tag} {cost}", "description": "d", "price": price, "cost_price": cost,
            "image_url": "https://e.com/a.png", "stock": stock, "branch_id": branch}, timeout=20)
        assert r.status_code == 200, r.text
        pids.append(r.json()["id"])
    try:
        r = director.get(f"{BASE_URL}/api/reports/inventory-value", params={"branch_ids": f"{b1},{b2}"}, timeout=20)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["source"] == "live"
        rows = {b["branch_id"]: b for b in data["branches"]}
        assert set(rows) == {b1, b2}
        assert (rows[b1]["cost_total"], rows[b1]["retail_total"], rows[b1]["profit"]) == (1000, 1800, 800)
        assert (rows[b2]["cost_total"], rows[b2]["retail_total"], rows[b2]["sku_count"]) == (500, 860, 1)
        assert data["total"]["profit"] == 1160
        assert len(data["items"]) == 2

        one = director.get(f"{BASE_URL}/api/reports/inventory-value", params={"branch_ids": b1}, timeout=20).json()
        assert [b["branch_id"] for b in one["branches"]] == [b1]

        assert director.get(f"{BASE_URL}/api/reports/inventory-value",
                            params={"date": "2000-01-01"}, timeout=20).status_code == 400
        assert director.get(f"{BASE_URL}/api/reports/inventory-value",
                            params={"date": "2999-01-01"}, timeout=20).status_code == 400

        x = director.get(f"{BASE_URL}/api/reports/inventory-value.xlsx", params={"branch_ids": f"{b1},{b2}"}, timeout=20)
        assert x.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(x.content))
        assert wb.sheetnames == ["Filiallar", "Mahsulotlar"]
    finally:
        for pid in pids:
            director.delete(f"{BASE_URL}/api/products/{pid}", timeout=20)
        director.delete(f"{BASE_URL}/api/branches/{b1}", timeout=20)
        director.delete(f"{BASE_URL}/api/branches/{b2}", timeout=20)


def test_import_merges_batches_with_weighted_average_cost():
    from stock_import import parse_stock_workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Номер", "Код товара", "Наименование", "Производитель", "Кол-во", "Штучно", "Срок.год",
               "Цена Пок", "Цена Прод", "Цена Опт", "Сумма Пок", "Сумма Прод", "Сумма Опт", "Штрих код", "Штуки"])
    ws.append([1, 7, "Omega", "X", 2, 0, "01.03.2028", 200, 290, 0, 400, 580, 0, "BC7", 90])
    ws.append([2, 7, "Omega", "X", 3, 0, "01.01.2028", 210, 300, 0, 630, 900, 0, "BC7", 90])
    buf = io.BytesIO()
    wb.save(buf)
    [row] = parse_stock_workbook(buf.getvalue())
    assert row["stock"] == 5 and row["cost_price"] == 206.0  # (2×200 + 3×210) / 5
    assert row["stock"] * row["cost_price"] == 1030           # fayldagi "Сумма Пок" bilan bir xil
    assert row["price"] == 300 and row["expiry_date"] == "2028-01-01"
